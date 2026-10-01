"""
Smart Exam Grader - Excel Exporter
Membaca roster siswa dari file Excel master, mencari baris siswa berdasarkan No Peserta / Nama,
dan menginjeksi hasil penilaian secara otomatis dengan styling profesional.
"""

import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from typing import Dict, Any, List, Optional
from thefuzz import process, fuzz

class ExcelExporter:
    def __init__(self, template_path: str = "data/templates/daftar_nilai_template.xlsx"):
        self.template_path = template_path

    def get_student_roster(self) -> List[Dict[str, Any]]:
        """Membaca daftar nomor peserta dan nama siswa dari template Excel."""
        if not os.path.exists(self.template_path):
            raise FileNotFoundError(f"Template Excel tidak ditemukan: {self.template_path}")

        wb = openpyxl.load_workbook(self.template_path, data_only=True)
        ws = wb.active

        roster = []
        # Header ada di baris 5, data mulai baris 6
        for row in range(6, ws.max_row + 1):
            no_val = ws.cell(row=row, column=1).value
            no_peserta = str(ws.cell(row=row, column=2).value or "").strip()
            nama = str(ws.cell(row=row, column=3).value or "").strip()

            if no_peserta and nama:
                roster.append({
                    "row_idx": row,
                    "no": no_val,
                    "no_peserta": no_peserta,
                    "nama": nama
                })

        wb.close()
        return roster

    def export_grade(self,
                     identity: Dict[str, Any],
                     score_data: Dict[str, Any],
                     output_path: str = "data/outputs/daftar_nilai_terkoreksi.xlsx") -> Dict[str, Any]:
        """
        Mengisi baris siswa di file Excel dengan hasil penilaian.
        Jika file output belum ada, copy dari template.
        """
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        if not os.path.exists(output_path):
            shutil.copyfile(self.template_path, output_path)

        wb = openpyxl.load_workbook(output_path)
        ws = wb.active

        # Cari baris siswa yang cocok
        target_row = None
        matched_by = None
        student_no = identity.get("no_peserta", "")
        student_name = identity.get("nama", "")

        # 1. Coba cocokkan via No Peserta
        if student_no:
            for row in range(6, ws.max_row + 1):
                cell_no = str(ws.cell(row=row, column=2).value or "").strip()
                if cell_no == student_no:
                    target_row = row
                    matched_by = "No Peserta"
                    break

        # 2. Jika belum cocok, coba cocokkan via Nama (Fuzzy Matching)
        if target_row is None and student_name:
            roster_names = {}
            for row in range(6, ws.max_row + 1):
                cell_name = str(ws.cell(row=row, column=3).value or "").strip()
                if cell_name:
                    roster_names[cell_name] = row

            if roster_names:
                best_match, score = process.extractOne(student_name, list(roster_names.keys()), scorer=fuzz.token_sort_ratio)
                if score >= 65:
                    target_row = roster_names[best_match]
                    matched_by = f"Nama Fuzzy Match ({best_match}, skor: {score}%)"

        if target_row is None:
            # Jika siswa tidak ada di daftar, tambahkan baris baru di bawah
            target_row = ws.max_row + 1
            ws.cell(row=target_row, column=1, value=ws.max_row - 4).alignment = Alignment(horizontal="center")
            ws.cell(row=target_row, column=2, value=student_no).alignment = Alignment(horizontal="center")
            ws.cell(row=target_row, column=3, value=student_name)
            matched_by = "Baris Baru (Tidak ditemukan di master)"

        # Injeksi Data Penilaian ke Kolom:
        # Col 4: Skor PG (25)
        # Col 5: Skor PGK (10)
        # Col 6: Total Skor
        # Col 7: Nilai Akhir (100)
        # Col 8: Status (LULUS / REMIDI)
        # Col 9: Catatan

        ws.cell(row=target_row, column=4, value=score_data["pg_correct"]).alignment = Alignment(horizontal="center")
        ws.cell(row=target_row, column=5, value=score_data["pgk_points"]).alignment = Alignment(horizontal="center")
        ws.cell(row=target_row, column=6, value=score_data["total_score_raw"]).alignment = Alignment(horizontal="center")

        grade_cell = ws.cell(row=target_row, column=7, value=score_data["final_grade_100"])
        grade_cell.alignment = Alignment(horizontal="center")
        grade_cell.font = Font(name="Calibri", size=11, bold=True)

        status_cell = ws.cell(row=target_row, column=8, value=score_data["status"])
        status_cell.alignment = Alignment(horizontal="center")

        if score_data["status"] == "LULUS":
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="166534") # Green
            status_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        else:
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="991B1B") # Red
            status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        catatan_cell = ws.cell(row=target_row, column=9, value=score_data["catatan"])
        if score_data.get("review_flags"):
            catatan_cell.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # Warning Yellow

        wb.save(output_path)
        wb.close()

        return {
            "success": True,
            "row_updated": target_row,
            "matched_by": matched_by,
            "output_file": output_path
        }
